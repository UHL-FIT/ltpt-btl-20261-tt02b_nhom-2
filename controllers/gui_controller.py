import tkinter as tk
from tkinter import ttk, messagebox, filedialog
import pandas as pd
from models import tinhhocbong as diemdanh
import views.gui_view as gui_view
from utils.logger import setup_logger

logger = setup_logger()

# Các biến toàn cục (module-level state)
app_df = pd.DataFrame()
app_ui = {}
app_root = None
app_edit_widget = None

# Ánh xạ: tên cột hiển thị → tên cột trong DataFrame
_COL_MAP = {
    "CC (0-10)":  "diem_cc",
    "GK (0-10)":  "diem_gk",
    "CK (0-10)":  "diem_ck",
    "RL (0-100)": "diem_rl",
}


def _tai_du_lieu():
    """Tải và hiển thị danh sách sinh viên lên bảng (kết hợp bộ lọc nếu có)."""
    global app_df
    app_df, ok = diemdanh.lay_danh_sach()
    if not ok:
        messagebox.showerror("Lỗi", "Không thể tải dữ liệu.")
        return

    search_text = app_ui['ent_search'].get().strip().lower()
    search_by   = app_ui['cbo_search_by'].get()

    display_df = app_df.copy()
    if search_text and not display_df.empty:
        if search_by == "MSV":
            display_df = display_df[display_df['msv'].astype(str).str.lower().str.contains(search_text, na=False)]
        elif search_by == "Họ Tên":
            display_df = display_df[display_df['ho_ten'].astype(str).str.lower().str.contains(search_text, na=False)]
        elif search_by == "Giới tính":
            display_df = display_df[display_df['gioi_tinh'].astype(str).str.lower().str.contains(search_text, na=False)]
        elif search_by == "SĐT":
            display_df = display_df[display_df['sdt'].astype(str).str.lower().str.contains(search_text, na=False)]
        elif search_by == "Xếp loại":
            display_df = display_df[display_df['xep_loai'].astype(str).str.lower().str.contains(search_text, na=False)]
        elif search_by == "Học bổng":
            display_df = display_df[display_df['du_hb'].astype(str).str.lower().str.contains(search_text, na=False)]
        else:  # Tất cả
            mask = display_df.apply(lambda row: row.astype(str).str.lower().str.contains(search_text).any(), axis=1)
            display_df = display_df[mask]

    # Sắp tên theo 'ho_ten' (A → Z) trước khi hiển thị
    if not display_df.empty and 'ho_ten' in display_df.columns:
        display_df['_sort_name'] = display_df['ho_ten'].fillna('').astype(str).str.lower()
        display_df = display_df.sort_values(by='_sort_name').drop(columns=['_sort_name'])

    gui_view.hien_thi_bang(app_ui, display_df)

    stats = diemdanh.thong_ke(display_df)
    gui_view.cap_nhat_thong_ke(app_ui, stats)


def on_search():
    """Xử lý sự kiện click nút Tìm kiếm."""
    logger.info("Người dùng thực hiện tìm kiếm.")
    _tai_du_lieu()


def on_clear_search():
    """Xóa trạng thái tìm kiếm và hiển thị lại toàn bộ danh sách."""
    logger.info("Người dùng xóa bộ lọc tìm kiếm.")
    app_ui['ent_search'].delete(0, tk.END)
    app_ui['cbo_search_by'].set("Tất cả")
    _tai_du_lieu()

def on_loc_hoc_bong():
    """Tự động điền điều kiện và lọc danh sách sinh viên đủ điều kiện học bổng."""
    logger.info("Người dùng click Lọc nhanh sinh viên đạt học bổng.")
    # Đặt tiêu chí tìm kiếm vào cột Học bổng và từ khóa là "Có" (tương ứng với "✅ Có")
    app_ui['cbo_search_by'].set("Học bổng")
    app_ui['ent_search'].delete(0, tk.END)
    app_ui['ent_search'].insert(0, "Có")
    _tai_du_lieu()


def on_nhap_diem():
    """Mở popup nhập điểm (CC, GK, CK, RL) cho sinh viên được chọn."""
    logger.info("Người dùng click Nhập Điểm.")
    global app_df
    tree = app_ui['tree']

    # Ưu tiên lấy SV được tick checkbox
    selected = [item_id for item_id in tree.get_children()
                if tree.item(item_id, 'values')[0] == "☑"]

    # Nếu không có, lấy dòng đang bôi xanh
    if not selected:
        selected = list(tree.selection())

    if not selected:
        messagebox.showwarning(
            "Cảnh báo",
            "Vui lòng tick chọn (☑) hoặc bấm chọn 1 sinh viên để nhập điểm!"
        )
        return

    if len(selected) > 1:
        messagebox.showwarning(
            "Cảnh báo",
            "Vui lòng chỉ chọn 1 sinh viên để nhập điểm!"
        )
        return

    item = tree.item(selected[0])
    msv    = str(item['values'][2]).strip().upper()
    ho_ten = item['values'][3]

    # Lấy điểm hiện tại từ DataFrame
    row_df = app_df[app_df['msv'].fillna('').astype(str).str.strip().str.upper() == msv]
    current_scores = {}
    if not row_df.empty:
        current_scores = {
            'diem_cc': float(row_df.iloc[0].get('diem_cc', 0)),
            'diem_gk': float(row_df.iloc[0].get('diem_gk', 0)),
            'diem_ck': float(row_df.iloc[0].get('diem_ck', 0)),
            'diem_rl': float(row_df.iloc[0].get('diem_rl', 0)),
        }

    scores = gui_view.hien_thi_form_nhap_diem(
        app_root, msv, ho_ten, current_scores
    )
    if scores:
        for col, val in scores.items():
            app_df, _ = diemdanh.cap_nhat_diem(app_df, msv, col, val)
        _tai_du_lieu()


def _select_student_by_msv(msv):
    tree = app_ui.get("tree")
    if not tree:
        return
    target_msv = str(msv).strip().upper()
    for item_id in tree.get_children():
        values = tree.item(item_id, "values")
        if str(values[2]).strip().upper() == target_msv:
            tree.selection_set(item_id)
            tree.see(item_id)
            break


def on_them_sv():
    """Bật cửa sổ Pop-up để thêm mới 1 sinh viên."""
    logger.info("Người dùng click Thêm Sinh viên.")
    global app_df
    data = gui_view.hien_thi_form_sinh_vien(app_root, is_edit=False)
    if data:
        app_df, ok, msg = diemdanh.them_sinh_vien(app_df, data)
        if ok:
            _tai_du_lieu()
            _select_student_by_msv(data.get("msv", ""))
        else:
            messagebox.showerror("Lỗi", msg)


def on_sua_sv():
    """Bật cửa sổ Pop-up để sửa thông tin sinh viên được tick chọn."""
    logger.info("Người dùng click Sửa Sinh viên.")
    global app_df
    tree = app_ui['tree']

    selected = []
    for item_id in tree.get_children():
        values = tree.item(item_id, 'values')
        if values[0] == "☑":
            selected.append(item_id)

    if not selected:
        selected = list(tree.selection())

    if not selected:
        messagebox.showwarning("Cảnh báo", "Vui lòng tick chọn (☑) hoặc bấm chọn 1 sinh viên để sửa!")
        return

    if len(selected) > 1:
        messagebox.showwarning("Cảnh báo", "Bạn đang chọn nhiều hơn 1 sinh viên. Vui lòng chỉ chọn 1 để sửa thông tin!")
        return

    item = tree.item(selected[0])
    # cols: Chọn(0) STT(1) MSV(2) Họ Tên(3) Giới tính(4) Lớp(5) SĐT(6)
    msv       = str(item['values'][2]).strip().upper()
    hoten     = item['values'][3]
    gioi_tinh = item['values'][4]
    lop       = item['values'][5]
    sdt       = item['values'][6]

    current_data = {"msv": msv, "ho_ten": hoten, "gioi_tinh": gioi_tinh, "lop": lop, "sdt": sdt}

    data = gui_view.hien_thi_form_sinh_vien(app_root, is_edit=True, current_data=current_data)
    if data:
        app_df, ok, msg = diemdanh.sua_sinh_vien(app_df, msv, data)
        if ok:
            _tai_du_lieu()
        else:
            messagebox.showerror("Lỗi", msg)


def on_xoa_sv():
    """Xóa các sinh viên được tick chọn khỏi danh sách."""
    logger.info("Người dùng click Xóa Sinh viên.")
    global app_df
    tree = app_ui['tree']
    msv_to_delete = []
    for item_id in tree.get_children():
        values = tree.item(item_id, 'values')
        if values[0] == "☑":
            msv_to_delete.append(values[2])

    if not msv_to_delete:
        messagebox.showwarning("Cảnh báo", "Vui lòng tick chọn (☑) ít nhất 1 sinh viên ở cột 'Chọn' để xóa!")
        return

    if messagebox.askyesno("Xác nhận", f"Bạn có chắc muốn xóa {len(msv_to_delete)} sinh viên đã chọn?"):
        app_df, ok, msg = diemdanh.xoa_nhieu_sinh_vien(app_df, msv_to_delete)
        if ok:
            _tai_du_lieu()
        else:
            messagebox.showerror("Lỗi", msg)


def on_import():
    """Import danh sách sinh viên từ một file Excel .xlsx."""
    logger.info("Người dùng click Import Excel.")
    global app_df
    filepath = filedialog.askopenfilename(
        title="Chọn file danh sách sinh viên (.xlsx)",
        filetypes=[
            ("Excel Files", "*.xlsx"),
            ("All Files", "*.*")
        ]
    )
    if not filepath:
        return

    try:
        if not filepath.lower().endswith('.xlsx'):
            raise ValueError("Chỉ hỗ trợ file Excel .xlsx")

        # Đọc file Excel (mặc định từ sheet đầu tiên)
        df_import = pd.read_excel(filepath, dtype=str)
        logger.info(f"Đọc file Excel: {filepath}")

        # Chuẩn hóa tên cột nhập từ Excel sang tên cột nội bộ
        col_map = {
            "msv": "msv",
            "mã sv": "msv",
            "mã sinh viên": "msv",
            "họ tên": "ho_ten",
            "ho_ten": "ho_ten",
            "tên": "ho_ten",
            "giới tính": "gioi_tinh",
            "gioi_tinh": "gioi_tinh",
            "lớp": "lop",
            "lop": "lop",
            "sđt": "sdt",
            "sdt": "sdt",
            "điểm cc": "diem_cc",
            "diem cc": "diem_cc",
            "cc": "diem_cc",
            "điểm gk": "diem_gk",
            "diem gk": "diem_gk",
            "gk": "diem_gk",
            "điểm ck": "diem_ck",
            "diem ck": "diem_ck",
            "ck": "diem_ck",
            "điểm rl": "diem_rl",
            "diem rl": "diem_rl",
            "rl": "diem_rl",
            "điểm tb": "diem_tb",
            "diem tb": "diem_tb",
            "xếp loại": "xep_loai",
            "xep loai": "xep_loai",
            "du_hb": "du_hb",
            "học bổng": "du_hb",
            "hoc bong": "du_hb",
        }

        df_import.columns = [
            col_map.get(str(col).strip().lower(), str(col).strip().lower())
            for col in df_import.columns
        ]

        if "ho_ten" not in df_import.columns or "msv" not in df_import.columns:
            messagebox.showerror(
                "Lỗi",
                "File phải có ít nhất hai cột: 'msv' và 'ho_ten'.\n\n"
                "Các cột cơ bản cần có:\n"
                "- msv, ho_ten\n"
                "- gioi_tinh, lop, sdt, diem_cc, diem_gk, diem_ck, diem_rl (tùy chọn)"
            )
            return

        def _safe_float(value):
            try:
                return float(value)
            except Exception:
                return 0.0

        existing_msv = set(app_df["msv"].astype(str).str.strip().str.upper())
        import_count = 0

        for _, row in df_import.iterrows():
            msv = str(row.get("msv", "")).strip().upper()
            ho_ten = str(row.get("ho_ten", "")).strip()
            if not msv or not ho_ten:
                continue

            if msv in existing_msv:
                continue

            new_row = {
                "msv": msv,
                "ho_ten": ho_ten,
                "gioi_tinh": str(row.get("gioi_tinh", "")).strip() or "Nam",
                "lop": str(row.get("lop", "")).strip(),
                "sdt": str(row.get("sdt", "")).strip(),
                "diem_cc": _safe_float(row.get("diem_cc", 0)),
                "diem_gk": _safe_float(row.get("diem_gk", 0)),
                "diem_ck": _safe_float(row.get("diem_ck", 0)),
                "diem_rl": _safe_float(row.get("diem_rl", 0)),
            }

            app_df = pd.concat([app_df, pd.DataFrame([new_row])], ignore_index=True)
            existing_msv.add(msv)
            import_count += 1

        if import_count:
            diemdanh.luu_danh_sach(app_df)
            _tai_du_lieu()

        messagebox.showinfo(
            "Thành công",
            f"Đã import thành công {import_count} sinh viên từ:\n{filepath}"
        )
        logger.info(f"Import thành công {import_count} sinh viên từ: {filepath}")
    except Exception as e:
        messagebox.showerror("Lỗi", f"Không thể import file: {e}")
        logger.error(f"Lỗi khi import: {e}", exc_info=True)


def on_export():
    """Xuất danh sách sinh viên hiện tại ra file Excel."""
    
    logger.info("Người dùng click Export Excel.")

    if app_df.empty:
        messagebox.showwarning(
            "Cảnh báo",
            "Không có dữ liệu để xuất!"
        )
        return

    filepath = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel Files", "*.xlsx")],
        title="Chọn nơi lưu file Excel"
    )

    if not filepath:
        return

    try:

        search_text = app_ui['ent_search'].get().strip().lower()
        search_by   = app_ui['cbo_search_by'].get()

        display_df = app_df.copy()

        if search_text and not display_df.empty:

            if search_by == "MSV":
                display_df = display_df[
                    display_df['msv'].astype(str)
                    .str.lower()
                    .str.contains(search_text)
                ]

            elif search_by == "Họ Tên":
                display_df = display_df[
                    display_df['ho_ten'].astype(str)
                    .str.lower()
                    .str.contains(search_text)
                ]

            elif search_by == "Giới tính":
                display_df = display_df[
                    display_df['gioi_tinh'].astype(str)
                    .str.lower()
                    .str.contains(search_text)
                ]

            elif search_by == "SĐT":
                display_df = display_df[
                    display_df['sdt'].astype(str)
                    .str.lower()
                    .str.contains(search_text)
                ]

            elif search_by == "Xếp loại":
                display_df = display_df[
                    display_df['xep_loai'].astype(str)
                    .str.lower()
                    .str.contains(search_text)
                ]

            elif search_by == "Học bổng":
                display_df = display_df[
                    display_df['du_hb'].astype(str)
                    .str.lower()
                    .str.contains(search_text)
                ]

            elif search_by == "Tất cả":

                mask = display_df.astype(str).apply(
                    lambda x: x.str.lower().str.contains(search_text)
                ).any(axis=1)

                display_df = display_df[mask]

        # ===== IMPORT OPENPYXL =====
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

        # ===== Đổi tên cột =====
        export_df = display_df.rename(columns={
            "msv": "Mã SV",
            "ho_ten": "Họ tên",
            "lop": "Lớp",
            "sdt": "SĐT",
            "gioi_tinh": "Giới tính",
            "diem_cc": "Điểm CC",
            "diem_gk": "Điểm GK",
            "diem_ck": "Điểm CK",
            "diem_rl": "Điểm RL",
            "diem_tb": "Điểm TB",
            "xep_loai": "Xếp loại",
            "du_hb": "Học bổng",
        })

        # ===== Tạo workbook =====
        wb = Workbook()
        ws = wb.active
        ws.title = "Danh sách sinh viên"

        # ===== Tiêu đề =====
        ws.merge_cells("A1:M1")

        ws["A1"] = "DANH SÁCH SINH VIÊN"

        ws["A1"].font = Font(
            bold=True,
            size=14
        )

        ws["A1"].alignment = Alignment(
            horizontal="center"
        )


        headers = [
            "STT",
            "MSV",
            "Họ tên",
            "Giới tính",
            "Lớp",
            "SĐT",
            "Điểm CC",
            "Điểm GK",
            "Điểm CK",
            "Điểm RL",
            "Điểm TB",
            "Xếp loại",
            "Học bổng"
        ]

        for col, header in enumerate(headers, start=1):

            cell = ws.cell(row=2, column=col)

            cell.value = header

            cell.font = Font(bold=True)

            cell.fill = PatternFill(
                start_color="FFF2CC",
                end_color="FFF2CC",
                fill_type="solid"
            )

            cell.alignment = Alignment(
                horizontal="center"
            )

        # ===== Dữ liệu =====
        for i, row in enumerate(export_df.values, start=3):

            ws.cell(i, 1, i - 2)

            for j, value in enumerate(row, start=2):

                ws.cell(i, j, value)

        # ===== Border =====
        thin = Side(style="thin")

        for row in ws.iter_rows(
            min_row=2,
            max_row=ws.max_row,
            min_col=1,
            max_col=13
        ):

            for cell in row:

                cell.border = Border(
                    left=thin,
                    right=thin,
                    top=thin,
                    bottom=thin
                )

        # ===== Độ 
       # Tìm đoạn cấu hình độ rộng cột này ở cuối hàm on_export trong file gui_controller.py
        # ===== Độ rộng cột =====
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 25
        ws.column_dimensions["C"].width = 25
        ws.column_dimensions["D"].width = 15
        ws.column_dimensions["E"].width = 15
        ws.column_dimensions["F"].width = 17
        ws.column_dimensions["G"].width = 12
        ws.column_dimensions["H"].width = 12
        ws.column_dimensions["I"].width = 12
        ws.column_dimensions["J"].width = 12
        ws.column_dimensions["K"].width = 15
        ws.column_dimensions["L"].width = 15
        ws.column_dimensions["M"].width = 18 # THÊM DÒNG NÀY: Để đặt độ rộng cho cột M (Học bổng) trong Excel
        # ===== Lưu file =====
        wb.save(filepath)

        messagebox.showinfo(
            "Thành công",
            f"Đã xuất file Excel:\n{filepath}"
        )

        logger.info(f"Xuất file Excel thành công: {filepath}")

    except Exception as e:

        messagebox.showerror(
            "Lỗi",
            f"Không thể xuất file:\n{e}"
        )

        logger.error(
            f"Lỗi khi export Excel: {e}",
            exc_info=True
        )
        
def on_about():
    """Hiển thị thông tin giới thiệu phần mềm."""
    logger.info("Người dùng click Giới thiệu.")
    about_text = (
        "🎓 PHẦN MỀM: SCHOLARSCORE\n"
        "------------------------------------\n"
        "🔹 Phiên bản: 2.0.0\n"
        "🔹 Chức năng: Xét Điểm Học Bổng Sinh viên\n"
        "------------------------------------\n"
        "Công thức:\n"
        "  Điểm TB = CC×0.1 + GK×0.3 + CK×0.6\n"
        "  Học Bổng: TB ≥ 8.0  VÀ  Rèn Luyện ≥ 80\n"
        "------------------------------------\n"
        "Xếp loại:\n"
        "  Xuất sắc: TB ≥ 9.0\n"
        "  Giỏi    : TB ≥ 8.0\n"
        "  Khá     : TB ≥ 6.5\n"
        "  Trung bình: TB ≥ 5.0\n"
        "  Yếu    : TB < 5.0"
    )
    messagebox.showinfo("Giới thiệu", about_text)


def on_single_click(event):
    tree = app_ui['tree']
    region     = tree.identify_region(event.x, event.y)
    column_str = tree.identify_column(event.x)

    if not column_str:
        return

    col_idx  = int(column_str.replace("#", "")) - 1
    col_name = app_ui['cols'][col_idx]

    # Click vào heading cột "Chọn" → chọn/bỏ chọn tất cả
    if region == "heading" and col_name == "Chọn":
        current_heading = tree.heading("Chọn", "text")
        if "☐" in current_heading:
            new_heading, new_val = "☑", "☑"
        else:
            new_heading, new_val = "☐", "☐"

        tree.heading("Chọn", text=new_heading)
        for item_id in tree.get_children():
            values = list(tree.item(item_id, "values"))
            values[0] = new_val
            tree.item(item_id, values=values)
        return

    # Click vào ô cột "Chọn" → toggle checkbox
    if region == "cell" and col_name == "Chọn":
        item_id = tree.identify_row(event.y)
        if item_id:
            values = list(tree.item(item_id, 'values'))
            values[0] = "☑" if values[0] == "☐" else "☐"
            tree.item(item_id, values=values)


def on_double_click(event):
    """Double-click vào ô điểm để chỉnh sửa trực tiếp."""
    global app_edit_widget, app_df
    if app_edit_widget:
        app_edit_widget.destroy()
        app_edit_widget = None

    tree = app_ui['tree']
    region = tree.identify_region(event.x, event.y)
    if region != "cell":
        return

    column_str = tree.identify_column(event.x)
    item_id    = tree.identify_row(event.y)

    if not item_id or not column_str:
        return

    col_idx  = int(column_str.replace("#", "")) - 1
    col_name = app_ui['cols'][col_idx]

    # Chỉ cho phép chỉnh sửa cột điểm
    if col_name not in _COL_MAP:
        return

    values  = tree.item(item_id, 'values')
    msv     = values[2]
    cur_val = str(values[col_idx])
    df_col  = _COL_MAP[col_name]

    # Giới hạn giá trị tối đa theo loại điểm
    max_val = 100.0 if col_name == "RL (0-100)" else 10.0

    x, y, w, h = tree.bbox(item_id, column_str)
    ent = ttk.Entry(tree)
    ent.insert(0, cur_val)
    ent.place(x=x, y=y, width=w, height=h)
    ent.focus_set()
    ent.select_range(0, tk.END)

    def save_diem(e=None):
        global app_df
        val_str = ent.get().strip()
        ent.destroy()
        try:
            new_val = float(val_str)
            if 0 <= new_val <= max_val:
                app_df, ok = diemdanh.cap_nhat_diem(app_df, msv, df_col, new_val)
                if ok:
                    _tai_du_lieu()
            else:
                messagebox.showwarning("Lỗi", f"Giá trị phải trong khoảng 0 – {max_val:.0f}.")
        except ValueError:
            messagebox.showwarning("Lỗi", "Giá trị phải là số hợp lệ.")

    ent.bind("<Return>",   save_diem)
    ent.bind("<FocusOut>", lambda e: ent.destroy())
    app_edit_widget = ent

def on_loc_hoc_bong():
    """Tự động điền điều kiện và lọc danh sách sinh viên đủ điều kiện học bổng."""
    logger.info("Người dùng click Lọc nhanh sinh viên đạt học bổng.")
    app_ui['cbo_search_by'].set("Học bổng")
    app_ui['ent_search'].delete(0, tk.END)
    app_ui['ent_search'].insert(0, "Có") # Đảm bảo giữ nguyên là "Có" để khớp với bảng của bạn
    _tai_du_lieu()

def _bind_events():
    app_ui['btn_them'].config(command=on_them_sv)
    app_ui['btn_sua'].config(command=on_sua_sv)
    app_ui['btn_xoa'].config(command=on_xoa_sv)
    app_ui['btn_import'].config(command=on_import)
    app_ui['btn_export'].config(command=on_export)
    app_ui['btn_nhap_diem'].config(command=on_nhap_diem)
    app_ui['btn_about'].config(command=on_about)

    app_ui['btn_search'].config(command=on_search)
    app_ui['btn_clear_search'].config(command=on_clear_search)
    
    # --- KẾT NỐI NÚT LỌC HỌC BỔNG ---
    app_ui['btn_loc_hb'].config(command=on_loc_hoc_bong)
    
    app_ui['ent_search'].bind("<Return>", lambda e: on_search())

    tree = app_ui['tree']
    tree.bind("<Double-1>",      on_double_click)
    tree.bind("<ButtonRelease-1>", on_single_click)
    
    # THÊM DÒNG NÀY ĐỂ KẾT NỐI NÚT BẤM VỚI HÀM XỬ LÝ
    app_ui['btn_loc_hb'].config(command=on_loc_hoc_bong)
    
    app_ui['ent_search'].bind("<Return>", lambda e: on_search())


def chay_ung_dung():
    """Khởi chạy ứng dụng GUI."""
    global app_root, app_ui
    logger.info("Khởi động ứng dụng Xét Điểm Học Bổng (GUI)")
    app_root = tk.Tk()

    app_ui = gui_view.tao_giao_dien_chinh(app_root)
    _bind_events()
    _tai_du_lieu()

    app_root.mainloop()
    logger.info("Thoát ứng dụng (GUI)")
